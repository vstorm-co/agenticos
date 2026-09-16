"""The eight new chat attachment formats, each producing useful text or a clean None.

The office/OpenDocument fixtures are generated in-process by the same libraries that
read them (a fake proves nothing); the legacy `.xls` is a committed fixture, since
`xlrd` only reads; the Outlook `.msg` reader is driven through a fake OLE object,
since no Python library writes the OLE compound format. A malformed input is a caught
`None` throughout — the upload has already succeeded, so raising would lose the file
rather than the parse (#1591).
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from app.services import file_upload as fu
from app.services.file_upload import FileUploadService

pytestmark = pytest.mark.anyio

FIXTURES = Path(__file__).parent / "fixtures"


def _odt(paragraphs: list[str]) -> bytes:
    from odf.opendocument import OpenDocumentText
    from odf.text import P

    document = OpenDocumentText()
    for text in paragraphs:
        document.text.addElement(P(text=text))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _ods(sheet: str, rows: list[list[str]]) -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    document = OpenDocumentSpreadsheet()
    table = Table(name=sheet)
    for row in rows:
        table_row = TableRow()
        for value in row:
            cell = TableCell()
            cell.addElement(P(text=value))
            table_row.addElement(cell)
        table.addElement(table_row)
    document.spreadsheet.addElement(table)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _odp(text: str) -> bytes:
    from odf.draw import Frame, Page, TextBox
    from odf.opendocument import OpenDocumentPresentation
    from odf.style import MasterPage, PageLayout
    from odf.text import P

    document = OpenDocumentPresentation()
    layout = PageLayout(name="pl1")
    document.automaticstyles.addElement(layout)
    document.masterstyles.addElement(MasterPage(name="m1", pagelayoutname="pl1"))
    page = Page(masterpagename="m1")
    frame = Frame()
    box = TextBox()
    box.addElement(P(text=text))
    frame.addElement(box)
    page.addElement(frame)
    document.presentation.addElement(page)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Quarterly"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "item"
    table.cell(0, 1).text = "cost"
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class TestOpenDocument:
    def test_an_odt_yields_its_paragraphs(self):
        text = FileUploadService._parse_odt_content(_odt(["Hello world", "Second para"]))

        assert text == "Hello world\nSecond para"

    def test_an_ods_yields_tab_separated_sheets(self):
        text = FileUploadService._parse_ods_content(
            _ods("Prices", [["item", "cost"], ["kawa", "1,5"]])
        )

        assert text == "Sheet: Prices\nitem\tcost\nkawa\t1,5"

    def test_an_odp_yields_its_slide_text(self):
        assert FileUploadService._parse_odp_content(_odp("Slide text here")) == "Slide text here"

    @pytest.mark.parametrize(
        "parser",
        [
            FileUploadService._parse_odt_content,
            FileUploadService._parse_ods_content,
            FileUploadService._parse_odp_content,
        ],
    )
    def test_malformed_opendocument_is_none(self, parser):
        assert parser(b"not an opendocument file at all") is None


class TestPresentation:
    def test_a_pptx_yields_shape_text_table_cells_and_notes(self):
        text = FileUploadService._parse_pptx_content(_pptx())

        assert text is not None
        assert "Quarterly" in text  # shape text
        assert "item\tcost" in text  # table cells
        assert "Notes: Speaker note" in text  # slide notes

    def test_malformed_pptx_is_none(self):
        assert FileUploadService._parse_pptx_content(b"still not a zip") is None


class TestLegacyXls:
    def test_a_committed_xls_reads_like_a_workbook(self):
        data = (FIXTURES / "legacy.xls").read_bytes()

        text = FileUploadService._parse_xls_content(data)

        assert text is not None
        assert "Sheet: Cover" in text
        assert "item\tcount\tprice\twhen" in text
        assert "kawa\t42\t1.5\t2024-03-01" in text  # int, float and ISO date

    def test_something_that_is_not_a_workbook_is_none(self):
        assert FileUploadService._parse_xls_content(b"not a workbook") is None


class _Stream:
    def __init__(self, data: bytes) -> None:
        self._data = data

    def read(self) -> bytes:
        return self._data


class _Ole:
    def __init__(self, streams: dict[str, bytes], dirs: list[list[str]]) -> None:
        self._streams = streams
        self._dirs = dirs
        self.closed = False

    def exists(self, name: str) -> bool:
        return name in self._streams

    def openstream(self, name: str) -> _Stream:
        return _Stream(self._streams[name])

    def listdir(self) -> list[list[str]]:
        return self._dirs

    def close(self) -> None:
        self.closed = True


def _u16(text: str) -> bytes:
    return text.encode("utf-16-le")


class TestOutlookMsg:
    def test_headers_body_recipients_and_attachment_names(self):
        ole = _Ole(
            {
                "__substg1.0_0037001F": _u16("Quarterly numbers"),
                "__substg1.0_0C1A001F": _u16("Alice"),
                "__substg1.0_1000001F": _u16("The body text."),
                "__recip_version1.0_#00000000/__substg1.0_3001001F": _u16("Bob"),
                "__recip_version1.0_#00000000/__substg1.0_39FE001F": _u16("bob@example.com"),
                "__attach_version1.0_#00000000/__substg1.0_3707001F": _u16("invoice.pdf"),
            },
            [
                ["__substg1.0_0037001F"],
                ["__recip_version1.0_#00000000", "__substg1.0_3001001F"],
                ["__attach_version1.0_#00000000", "__substg1.0_3707001F"],
            ],
        )

        text = fu._msg_text(ole)

        assert text is not None
        assert "From: Alice" in text
        assert "To: Bob <bob@example.com>" in text
        assert "Subject: Quarterly numbers" in text
        assert "The body text." in text
        assert "Attachments: invoice.pdf" in text

    def test_the_html_body_is_the_fallback_stripped_of_tags(self):
        ole = _Ole({"__substg1.0_10130102": b"<html><body>Hello <b>bold</b></body></html>"}, [])

        assert fu._msg_text(ole) == "Hello bold"

    def test_mapi_nul_terminators_are_stripped(self):
        """`PtypString` streams are NUL-terminated and `str.strip()` leaves `\\x00`,
        which PostgreSQL's `Text` column rejects — failing the upload after the bytes
        were stored. Every string property drops it (#1591)."""
        ole = _Ole(
            {
                "__substg1.0_0037001F": _u16("Subject\x00"),
                "__substg1.0_1000001F": _u16("Body\x00text\x00"),
                "__recip_version1.0_#00000000/__substg1.0_3001001F": _u16("Bob\x00"),
                "__recip_version1.0_#00000000/__substg1.0_39FE001F": _u16("bob@x.com\x00"),
                "__attach_version1.0_#00000000/__substg1.0_3707001F": _u16("a.pdf\x00"),
            },
            [
                ["__substg1.0_0037001F"],
                ["__recip_version1.0_#00000000", "__substg1.0_3001001F"],
                ["__attach_version1.0_#00000000", "__substg1.0_3707001F"],
            ],
        )

        text = fu._msg_text(ole)

        assert text is not None
        assert "\x00" not in text
        assert "Subject: Subject" in text
        assert "Bodytext" in text
        assert "To: Bob <bob@x.com>" in text
        assert "Attachments: a.pdf" in text

    def test_the_html_body_nul_is_stripped(self):
        ole = _Ole({"__substg1.0_10130102": b"<p>Hi\x00 there</p>"}, [])

        result = fu._msg_text(ole)

        assert result is not None
        assert "\x00" not in result
        assert result == "Hi there"

    def test_an_ansi_property_is_read_and_stripped(self):
        ole = _Ole(
            {"__substg1.0_0037001E": "Legacy\x00".encode("cp1252")},
            [["__substg1.0_0037001E"]],
        )

        assert fu._msg_text(ole) == "Subject: Legacy"

    def test_a_non_ole_file_is_none(self):
        assert FileUploadService._parse_msg_content(b"just some plain text") is None

    def test_the_ole_stream_is_read_and_closed(self, monkeypatch):
        ole = _Ole({"__substg1.0_0037001F": _u16("Hi")}, [["__substg1.0_0037001F"]])
        import olefile

        monkeypatch.setattr(olefile, "isOleFile", lambda _stream: True)
        monkeypatch.setattr(olefile, "OleFileIO", lambda _stream: ole)

        text = FileUploadService._parse_msg_content(b"\xd0\xcf\x11\xe0anything")

        assert text is not None
        assert "Subject: Hi" in text
        assert ole.closed is True


class TestDocViaLibreOffice:
    async def test_a_doc_is_extracted_through_the_managed_converter(self, monkeypatch):
        from unittest.mock import AsyncMock

        convert = AsyncMock(return_value="Dear all, the contract is agreed.")
        monkeypatch.setattr(fu, "libreoffice_convert", convert)

        text = await FileUploadService(db=None)._parse_doc_content(b"\xd0\xcf\x11\xe0doc")

        assert text == "Dear all, the contract is agreed."
        assert convert.await_args.kwargs["suffix"] == ".doc"

    async def test_an_absent_converter_degrades_to_none(self, monkeypatch):
        from unittest.mock import AsyncMock

        monkeypatch.setattr(fu, "libreoffice_convert", AsyncMock(return_value=None))

        assert await FileUploadService(db=None)._parse_doc_content(b"doc") is None

    async def test_a_converter_error_is_caught(self, monkeypatch):
        from unittest.mock import AsyncMock

        monkeypatch.setattr(fu, "libreoffice_convert", AsyncMock(side_effect=RuntimeError("boom")))

        assert await FileUploadService(db=None)._parse_doc_content(b"doc") is None


class TestTheTextBudget:
    def test_extracted_text_is_capped_with_a_marker(self):
        capped = fu.cap_text("A" * 100, 10)

        assert capped is not None
        assert capped.startswith("A" * 10)
        assert "[truncated 10 of 100 chars]" in capped

    def test_text_within_the_cap_is_untouched(self):
        assert fu.cap_text("short", 10) == "short"
        assert fu.cap_text(None, 10) is None


class TestXmlDecoding:
    def test_a_utf16_xml_document_decodes(self):
        document = '<?xml version="1.0" encoding="UTF-16"?><root>café</root>'
        data = b"\xff\xfe" + document.encode("utf-16-le")

        text = FileUploadService._parse_text_content(data)

        assert text is not None
        assert "café" in text

    def test_a_declared_encoding_without_a_bom_is_honoured(self):
        data = "<?xml version='1.0' encoding='latin-1'?><r>é</r>".encode("latin-1")

        text = FileUploadService._parse_text_content(data)

        assert text is not None and "é" in text

    def test_a_declaration_with_spaces_around_the_equals_is_honoured(self):
        # The XML `Eq` production allows whitespace: `encoding = "..."` is legal and
        # must not drop the parse to `None` (#1591 second-pass).
        data = "<?xml version='1.0' encoding = \"windows-1250\"?><r>ł</r>".encode("cp1250")

        text = FileUploadService._parse_text_content(data)

        assert text is not None and "ł" in text

    def test_undecodable_bytes_are_none(self):
        assert FileUploadService._parse_text_content(b"\xff\xfe\xff\xfe\x00") is None


class TestZipBackedFormatsRunThroughTheGuard:
    """DOCX and XLSX are ZIP+XML like ODF/PPTX, so they must pass through
    `safe_unzip` before their parser opens them; otherwise a decompression bomb
    reaches `python-docx`/`openpyxl` unbounded (#1591, §7 finding 1)."""

    def _docx(self, text: str) -> bytes:
        from docx import Document

        document = Document()
        document.add_paragraph(text)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _xlsx(self, value: str) -> bytes:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active["A1"] = value
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_a_valid_docx_still_parses(self):
        assert FileUploadService._parse_docx_content(self._docx("Hello world")) == "Hello world"

    def test_a_valid_xlsx_still_parses(self):
        text = FileUploadService._parse_spreadsheet_content(self._xlsx("hi"))
        assert text is not None and "hi" in text

    def test_a_docx_tripping_the_member_cap_is_refused(self, monkeypatch):
        from app.core import config as config_module

        docx = self._docx("Hello world")
        monkeypatch.setattr(config_module.settings, "CHAT_ARCHIVE_MEMBER_MAX_BYTES", 1)

        assert FileUploadService._parse_docx_content(docx) is None

    def test_an_xlsx_tripping_the_member_cap_is_refused(self, monkeypatch):
        from app.core import config as config_module

        xlsx = self._xlsx("hi")
        monkeypatch.setattr(config_module.settings, "CHAT_ARCHIVE_MEMBER_MAX_BYTES", 1)

        assert FileUploadService._parse_spreadsheet_content(xlsx) is None


class TestOdsRepetitionIsBounded:
    """A tiny ODS can declare a colossal `number-columns-repeated`; the parser must
    not allocate `[text] * repeat` unbounded (#1591, §7 finding 3)."""

    def _ods_repeated(self, repeat: int, *, value: str = "x") -> bytes:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        document = OpenDocumentSpreadsheet()
        table = Table(name="S")
        row = TableRow()
        cell = TableCell(numbercolumnsrepeated=str(repeat))
        cell.addElement(P(text=value))
        row.addElement(cell)
        table.addElement(row)
        document.spreadsheet.addElement(table)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_a_billion_column_repeat_does_not_allocate_a_billion_cells(self):
        # Returns bounded text rather than exhausting memory; the cap is far below a
        # billion, so the extraction is capped, not the process.
        text = FileUploadService._parse_ods_content(self._ods_repeated(1_000_000_000))

        assert text is not None
        # ~1M cells of "x" joined by tabs, not a billion: bounded by the cell budget
        # (the "Sheet: S" header and the absent trailing tab account for the slack).
        assert len(text) <= fu._ODS_MAX_CELLS * 2 + 100

    def test_a_large_cell_value_repeated_is_bounded_by_output_size_not_cell_count(
        self, monkeypatch
    ):
        """The cell budget bounds the *list* of references, but `"\\t".join(cells)`
        materialises the value once per reference: a large value with a huge repeat is
        a multi-gigabyte string the cell budget alone does not stop. The character
        budget bounds the join (#1591, §7 finding 3)."""
        from app.core import config as config_module

        monkeypatch.setattr(config_module.settings, "CHAT_PARSED_TEXT_MAX_CHARS", 100_000)
        # A 10 KB value with a million repeats would be ~10 GB under the cell budget
        # alone; the character budget caps the built string near the parsed-text cap.
        text = FileUploadService._parse_ods_content(
            self._ods_repeated(1_000_000, value="v" * 10_000)
        )

        assert text is not None
        assert len(text) <= 100_000 + 10_000 + 100


class TestArchiveBombGuards:
    def _zip(self, members: dict[str, bytes]) -> bytes:
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, payload in members.items():
                archive.writestr(name, payload)
        return buffer.getvalue()

    def test_a_valid_archive_returns_a_stream(self):
        stream = fu.safe_unzip(self._zip({"a.xml": b"<x/>"}))

        assert stream.read()  # a usable, non-empty stream over the bytes

    def test_too_many_members_is_refused(self, monkeypatch):
        from app.core import config as config_module

        monkeypatch.setattr(config_module.settings, "CHAT_ARCHIVE_MAX_MEMBERS", 1)

        with pytest.raises(ValueError, match="too many members"):
            fu.safe_unzip(self._zip({"a": b"1", "b": b"2"}))

    def test_a_member_over_the_cap_is_refused_by_its_real_size(self, monkeypatch):
        """The bound is measured by reading the member, not by trusting the
        forgeable central-directory size."""
        from app.core import config as config_module

        monkeypatch.setattr(config_module.settings, "CHAT_ARCHIVE_MEMBER_MAX_BYTES", 1000)

        with pytest.raises(ValueError, match="member too large"):
            fu.safe_unzip(self._zip({"big": b"\x00" * 100_000}))

    def test_the_running_total_is_capped(self, monkeypatch):
        from app.core import config as config_module

        monkeypatch.setattr(config_module.settings, "CHAT_ARCHIVE_TOTAL_MAX_BYTES", 1500)

        with pytest.raises(ValueError, match="archive too large"):
            fu.safe_unzip(self._zip({f"m{i}": b"\x00" * 1000 for i in range(5)}))


class TestOdfSpaceBombIsBounded:
    """odfpy's `extractText` expands `<text:s text:c=N>` to `" " * N` before any cap,
    so a tiny ODF can allocate gigabytes of spaces. Each run is clamped first
    (#1591, §7 finding 3, security follow-up)."""

    def _odt_with_spaces(self, count: int, tail: str = "END") -> bytes:
        from odf.opendocument import OpenDocumentText
        from odf.text import P, S

        document = OpenDocumentText()
        p = P()
        p.addElement(S(c=count))
        p.addText(tail)
        document.text.addElement(p)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _odp_with_spaces(self, count: int) -> bytes:
        from odf.draw import Frame, Page, TextBox
        from odf.opendocument import OpenDocumentPresentation
        from odf.style import MasterPage, PageLayout
        from odf.text import P, S

        document = OpenDocumentPresentation()
        layout = PageLayout(name="pl1")
        document.automaticstyles.addElement(layout)
        document.masterstyles.addElement(MasterPage(name="m1", pagelayoutname="pl1"))
        page = Page(masterpagename="m1")
        frame = Frame()
        box = TextBox()
        p = P()
        p.addElement(S(c=count))
        p.addText("SLIDE")
        box.addElement(p)
        frame.addElement(box)
        page.addElement(frame)
        document.presentation.addElement(page)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _ods_with_spaces(self, count: int) -> bytes:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P, S

        document = OpenDocumentSpreadsheet()
        table = Table(name="S")
        row = TableRow()
        cell = TableCell()
        p = P()
        p.addElement(S(c=count))
        p.addText("V")
        cell.addElement(p)
        row.addElement(cell)
        table.addElement(row)
        document.spreadsheet.addElement(table)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_an_odt_space_bomb_is_bounded(self):
        text = FileUploadService._parse_odt_content(self._odt_with_spaces(10_000_000_000))

        assert text is not None
        assert "END" in text
        assert len(text) <= fu.settings.CHAT_PARSED_TEXT_MAX_CHARS + 100

    def test_an_odp_space_bomb_is_bounded(self):
        text = FileUploadService._parse_odp_content(self._odp_with_spaces(10_000_000_000))

        assert text is not None
        assert "SLIDE" in text
        assert len(text) <= fu.settings.CHAT_PARSED_TEXT_MAX_CHARS + 100

    def test_an_ods_cell_space_bomb_is_bounded(self):
        text = FileUploadService._parse_ods_content(self._ods_with_spaces(10_000_000_000))

        assert text is not None
        assert "V" in text
        assert len(text) <= fu.settings.CHAT_PARSED_TEXT_MAX_CHARS + 100

    def test_a_modest_space_run_is_preserved(self):
        text = FileUploadService._parse_odt_content(self._odt_with_spaces(4))

        assert text == "    END"

    def test_a_space_element_without_a_count_is_one_space(self):
        from odf.opendocument import OpenDocumentText
        from odf.text import P, S

        document = OpenDocumentText()
        p = P()
        p.addText("A")
        p.addElement(S())  # a bare `<text:s/>`, no `text:c` — one space
        p.addText("B")
        document.text.addElement(p)
        buffer = io.BytesIO()
        document.save(buffer)

        text = FileUploadService._parse_odt_content(buffer.getvalue())

        assert text == "A B"

    def test_a_zero_count_space_run_clamps_to_nothing(self):
        text = FileUploadService._parse_odt_content(self._odt_with_spaces(0, tail="END"))

        assert text == "END"


class TestOdsRowRepetitionIsExpandedAndBounded:
    """A nonempty row carrying `table:number-rows-repeated` is real data repeated, so
    it must be emitted that many times — but the count is attacker-controlled, so the
    expansion is charged to the same budgets the columns answer to (#1591)."""

    def _ods_rows(self, repeat: int, *, value: str = "hi") -> bytes:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        document = OpenDocumentSpreadsheet()
        table = Table(name="S")
        row = TableRow(numberrowsrepeated=str(repeat))
        cell = TableCell()
        cell.addElement(P(text=value))
        row.addElement(cell)
        table.addElement(row)
        document.spreadsheet.addElement(table)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_a_repeated_nonempty_row_is_emitted_each_time(self):
        text = FileUploadService._parse_ods_content(self._ods_rows(3))

        assert text is not None
        assert text.count("hi") == 3

    def test_a_colossal_row_repeat_is_bounded(self, monkeypatch):
        from app.core import config as config_module

        monkeypatch.setattr(config_module.settings, "CHAT_PARSED_TEXT_MAX_CHARS", 10_000)

        text = FileUploadService._parse_ods_content(self._ods_rows(1_000_000_000, value="row"))

        assert text is not None
        # Bounded near the char budget rather than a billion rows; the newline joins
        # the budget does not count add a bounded per-row slack, so ~2x is the ceiling.
        assert len(text) <= 10_000 * 2


class TestLegacyXlsIsBounded:
    """A sparse BIFF sheet can describe a 65_536x256 rectangle for one value; the sweep
    is bounded by a cell budget so it cannot occupy the file pool (#1591)."""

    class _Cell:
        def __init__(self, value: str) -> None:
            import xlrd

            self.ctype = xlrd.XL_CELL_TEXT if value else xlrd.XL_CELL_EMPTY
            self.value = value

    class _WideSheet:
        """A rectangle every cell of which is populated, to exercise the budget."""

        def __init__(self, name: str, nrows: int, ncols: int) -> None:
            self.name = name
            self.nrows = nrows
            self.ncols = ncols
            self.reads = 0

        def row_len(self, _r: int) -> int:
            return self.ncols

        def cell(self, _r: int, _c: int) -> TestLegacyXlsIsBounded._Cell:
            self.reads += 1
            return TestLegacyXlsIsBounded._Cell("x")

    class _RowsSheet:
        """A sheet defined by explicit rows, `""` standing for a blank cell."""

        def __init__(self, name: str, rows: list[list[str]]) -> None:
            self.name = name
            self._rows = rows
            self.nrows = len(rows)
            self.ncols = max((len(r) for r in rows), default=0)

        def row_len(self, r: int) -> int:
            return len(self._rows[r])

        def cell(self, r: int, c: int) -> TestLegacyXlsIsBounded._Cell:
            return TestLegacyXlsIsBounded._Cell(self._rows[r][c])

    class _Book:
        def __init__(self, sheet: object) -> None:
            self._sheet = sheet
            self.datemode = 0

        def sheets(self) -> list[object]:
            return [self._sheet]

    def test_a_sparse_rectangle_does_not_sweep_every_cell(self, monkeypatch):
        import xlrd

        sheet = self._WideSheet("Big", nrows=65_536, ncols=256)
        monkeypatch.setattr(
            xlrd, "open_workbook", lambda **_kwargs: TestLegacyXlsIsBounded._Book(sheet)
        )

        text = FileUploadService._parse_xls_content(b"\xd0\xcf\x11\xe0anything")

        assert text is not None
        # Bounded near the cell budget rather than the full 16.8M-cell rectangle.
        assert sheet.reads <= fu._XLS_MAX_CELLS + 256

    def test_trailing_blanks_are_trimmed_and_empty_rows_dropped(self, monkeypatch):
        import xlrd

        sheet = self._RowsSheet("Cover", [["a", "b", "", ""], ["", "", ""], ["c", ""]])
        monkeypatch.setattr(
            xlrd, "open_workbook", lambda **_kwargs: TestLegacyXlsIsBounded._Book(sheet)
        )

        text = FileUploadService._parse_xls_content(b"\xd0\xcf\x11\xe0anything")

        assert text == "Sheet: Cover\na\tb\nc"

    def test_a_wholly_blank_sheet_yields_no_text(self, monkeypatch):
        import xlrd

        sheet = self._RowsSheet("Empty", [["", ""], [""]])
        monkeypatch.setattr(
            xlrd, "open_workbook", lambda **_kwargs: TestLegacyXlsIsBounded._Book(sheet)
        )

        assert FileUploadService._parse_xls_content(b"\xd0\xcf\x11\xe0anything") is None
